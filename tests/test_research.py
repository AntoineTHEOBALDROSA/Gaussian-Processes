import io
import json
import unittest
import warnings
import zipfile
from unittest.mock import patch
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment
from sklearn.gaussian_process.kernels import RBF, WhiteKernel, ConstantKernel
from acquisition import suggestions, latent_predictions
from excel_prediction import describe_workbook, inspect_sheet, predict_workbook
from models import GPRWrapper, noyaux_candidats, bound_diagnostics, search_options
from model_io import open_model, save_model
import test_exports
from test_exports import new_client


def example_workbook():
    book = Workbook(); sheet = book.active; sheet.title = 'Mesures'
    sheet.append(['Titre']); sheet.append([None,'X2','X','Référence','Autre modèle','FprédiX'])
    sheet.append([0,'0,7',.13,4.36,4.32,None])
    sheet.append([1,.8,.2,3.93,3.99,123.])
    sheet.append([2,None,.3,4.19,4.21,None])
    sheet.append([3,.9,.4,4.48,4.52,'=D6'])
    sheet['F3'].number_format = '0.0000'; sheet['F3'].fill=PatternFill('solid',fgColor='FFDDEE')
    sheet['F3'].comment = Comment('Cellule à compléter','Essai')
    sheet.freeze_panes = 'B3'; sheet.column_dimensions['F'].width = 25
    book.create_sheet('À conserver')['A1'] = '=Mesures!D3+1'
    output=io.BytesIO(); book.save(output); return output.getvalue()


class ResearchTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Reuse a trained, persisted fixture; the tested operations must never fit.
        test_exports.ExportTests.setUpClass()
        cls.client=new_client()
        response=cls.client.post('/api/models/open',data={'file':(io.BytesIO(test_exports.ExportTests.archive),'fixture.gpmodel')})
        cls.job=response.get_json(); cls.manifest,cls.csv,cls.model=open_model(io.BytesIO(test_exports.ExportTests.archive))
        cls.raw=example_workbook()
        cls.excel=dict(sheet='Mesures',header_row=2,first_row=3,last_row=6,mapping={'X':3,'X2':2},output_column=6)

    def test_excel_preserves_other_cells_parts_and_styles(self):
        info=inspect_sheet(self.raw,'Mesures',2)
        self.assertEqual(info['columns'][0]['name'],'')
        self.assertEqual(info['first_row'],3)
        with patch('models.GPRWrapper.fit',side_effect=AssertionError('No retraining')):
            raw,summary=predict_workbook(self.raw,self.model,['X','X2'],self.excel)
        self.assertEqual(summary,dict(predicted=1,skipped_invalid=1,skipped_existing=2,invalid_rows=[5],output='F'))
        book=load_workbook(io.BytesIO(raw)); ws=book['Mesures']
        self.assertAlmostEqual(ws['F3'].value,self.model.predict([[.13,.7]])[0])
        self.assertEqual(ws['F4'].value,123.)
        self.assertEqual(ws['F6'].value,'=D6')
        self.assertEqual(ws['F2'].value,'FprédiX')
        self.assertEqual(ws['F3'].number_format,'0.0000')
        self.assertEqual(ws['F3'].fill.fgColor.rgb,'00FFDDEE')
        self.assertEqual(ws['F3'].comment.text,'Cellule à compléter')
        self.assertEqual(ws.freeze_panes,'B3')
        self.assertEqual(book['À conserver']['A1'].value,'=Mesures!D3+1')
        with zipfile.ZipFile(io.BytesIO(self.raw)) as original,zipfile.ZipFile(io.BytesIO(raw)) as changed:
            for name in original.namelist():
                if name not in ('xl/workbook.xml','xl/worksheets/sheet1.xml'):
                    self.assertEqual(changed.read(name),original.read(name),name)

    def test_excel_new_column_overwrite_and_validation(self):
        raw,summary=predict_workbook(self.raw,self.model,['X','X2'],dict(self.excel,overwrite=True))
        self.assertEqual(summary['predicted'],3)
        ws=load_workbook(io.BytesIO(raw)).active
        self.assertAlmostEqual(ws['F6'].value,self.model.predict([[.4,.9]])[0])
        raw,summary=predict_workbook(self.raw,self.model,['X','X2'],dict(self.excel,output_column=7,output_name='F prédit'))
        ws=load_workbook(io.BytesIO(raw)).active
        self.assertEqual(ws['G2'].value,'F prédit');self.assertEqual(ws['F6'].value,'=D6')
        self.assertEqual(summary['predicted'],3)
        for change in ({'sheet':[]},{'mapping':{'X':3,'X2':3}},{'output_column':2},{'last_row':2},
                       {'first_row':True},{'last_row':10005},{'overwrite':'oui'}):
            with self.subTest(change=change),self.assertRaises(ValueError):
                predict_workbook(self.raw,self.model,['X','X2'],dict(self.excel,**change))
        with self.assertRaises(ValueError):describe_workbook(b'bad','bad.xlsx','id')

    def test_excel_api_with_reopened_model(self):
        client=self.client
        upload=client.post('/api/workbooks',data={'file':(io.BytesIO(self.raw),'Mesures.xlsx')})
        self.assertEqual(upload.status_code,200,upload.get_json()); workbook=upload.get_json()
        inspected=client.post('/api/workbooks/'+workbook['id']+'/inspect',json={'sheet':'Mesures','header_row':2})
        self.assertEqual(inspected.status_code,200)
        with patch('models.GPRWrapper.fit',side_effect=AssertionError('No fit')):
            response=client.post('/api/jobs/'+self.job['id']+'/predict-excel',json=dict(self.excel,workbook_id=workbook['id']))
        self.assertEqual(response.status_code,200,response.get_json())
        self.assertEqual(json.loads(response.headers['X-GP-Summary'])['predicted'],1)
        self.assertIsInstance(load_workbook(io.BytesIO(response.data)).active['F3'].value,float)
        self.assertEqual(client.post('/api/jobs/'+self.job['id']+'/predict-excel',json=dict(self.excel,workbook_id='absent')).status_code,400)

    def test_all_eight_kernels_save_without_changing_predictions(self):
        frame=test_exports.ExportTests.frame; train=self.manifest['result']['data']['train']
        X=frame[['X','X2']].to_numpy()[train];y=frame.F.to_numpy()[train]
        bank=noyaux_candidats(2,extended=True);self.assertEqual(len(bank),8)
        for name,kernel in bank.items():
            with self.subTest(kernel=name),warnings.catch_warnings():
                warnings.simplefilter('ignore')
                model=GPRWrapper(kernel,0).fit(X,y)
                fixture=dict(model=model,csv=self.csv,dataset=self.manifest['dataset'],options=self.manifest['options'],result=self.manifest['result'])
                _,_,saved=open_model(io.BytesIO(save_model(fixture)))
                points=np.array([[.12,.44],[3,10]])
                for before,after in zip(model.predict(points,return_std=True),saved.predict(points,return_std=True)):
                    np.testing.assert_allclose(before,after,atol=1e-8,rtol=1e-8)

    def test_bounds_diagnostics_and_options(self):
        kernel=ConstantKernel(1)*RBF([1000,1],(.01,1000))+WhiteKernel(1e-8,(1e-8,1))
        records=bound_diagnostics(kernel,['X','X2'])
        self.assertTrue(any(r['name']=='Longueur de X' and r['hit']=='upper' for r in records))
        self.assertTrue(any('déterministe' in r['explanation'] for r in records))
        self.assertEqual(search_options(dict(device='cpu'))[0],'base')
        for options in (dict(device='cuda',kernel_search='extended'),dict(length_bounds=[0,10]),dict(length_bounds=[True,10]),dict(length_bounds=[1,float('inf')])):
            with self.assertRaises(ValueError):search_options(options)

    def test_candidates_discrete_unique_in_range_reproducible_and_filtered(self):
        result=self.manifest['result']
        options=dict(count=10,domain=[dict(name='X',minimum=-3,maximum=3,integer=True),dict(name='X2',minimum=0,maximum=5,integer=True)])
        a=suggestions(self.model,result,options); b=suggestions(self.model,result,options)
        self.assertEqual(a,b)
        points=np.array([r['values'] for r in a['rows']]);known=np.array(result['data']['x'])
        self.assertEqual(len(np.unique(points,axis=0)),10)
        self.assertTrue((points==np.round(points)).all())
        self.assertTrue(((points[:,0]>=-3)&(points[:,0]<=3)&(points[:,1]>=0)&(points[:,1]<=5)).all())
        self.assertTrue(all(not np.any(np.all(known==p,axis=1)) for p in points))
        self.assertTrue(all(r['latent_std']<=r['std'] for r in a['rows']))
        means=[r['mean'] for r in a['rows']];low,high=min(means)-.01,max(means)+.01
        focused=suggestions(self.model,result,dict(options,target_range=[low,high]))
        self.assertTrue(all(low<=r['mean']<=high for r in focused['rows']))
        response=self.client.post('/api/jobs/'+self.job['id']+'/suggestions',json=options)
        self.assertEqual(response.status_code,200,response.get_json())
        self.assertEqual(response.get_json(),a)
        for change in (dict(count=0),dict(target_range=[1e10,1e11]),dict(target_range=[4,2])):
            with self.assertRaises(ValueError):suggestions(self.model,result,dict(options,**change))

    def test_noise_is_excluded_from_acquisition_uncertainty(self):
        X=np.linspace(-1,1,10)[:,None];y=np.sin(X[:,0])
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            model=GPRWrapper(RBF(1,'fixed')+WhiteKernel(.5,'fixed'),0).fit(X,y)
        _,std,latent=latent_predictions(model,np.array([[0.],[2.]]))
        np.testing.assert_allclose(std**2-latent**2,.5*model.y_std**2,atol=1e-12)

    def test_acquisition_identifies_a_hole_between_observed_clusters(self):
        X=np.r_[np.linspace(0,.1,10),np.linspace(.9,1,10)][:,None]
        y=np.sin(6*X[:,0])
        model=GPRWrapper(RBF(.2,'fixed')+WhiteKernel(1e-4,'fixed'),0).fit(X,y)
        result=dict(features=['X'],target='F',data=dict(x=X.tolist(),y=y.tolist()),
                    input_ranges=[dict(name='X',minimum=0,maximum=1)])
        options=dict(count=3,domain=[dict(name='X',minimum=0,maximum=1)])
        proposed=suggestions(model,result,options)
        self.assertTrue(.3<proposed['rows'][0]['values'][0]<.7)
        # A domain containing one existing configuration offers no new experiment.
        options['domain']=[dict(name='X',minimum=0,maximum=0)]
        with self.assertRaises(ValueError):suggestions(model,result,options)

    def test_excel_rejects_merged_output_cells(self):
        book=load_workbook(io.BytesIO(self.raw));book.active.merge_cells('F3:F4')
        raw=io.BytesIO();book.save(raw)
        with self.assertRaisesRegex(ValueError,'fusionnées'):
            predict_workbook(raw.getvalue(),self.model,['X','X2'],self.excel)


if __name__=='__main__':unittest.main()
